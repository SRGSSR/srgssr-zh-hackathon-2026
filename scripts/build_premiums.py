"""Build the health-insurance premium index (data/premiums_<year>.sqlite).

Sources (all Federal Office of Public Health, BAG):
  * Premium data  : opendata.swiss dataset "health-insurance-premiums" (opendata.bagnet.ch)
  * Premium regions: https://www.priminfo.admin.ch/downloads/praemienregionen.xlsx
  * Insurer names  : BAG "Verzeichnis der zugelassenen Krankenversicherer"

Usage:
    uv run --extra build python scripts/build_premiums.py --year 2026
"""

from __future__ import annotations

import argparse
import base64
import csv
import io
import sqlite3
import urllib.parse
import zipfile
from datetime import date
from pathlib import Path

import httpx
import openpyxl

UA = {"User-Agent": "swiss-grounding-mcp/0.1 (+hackathon prototype)"}
ROOT = Path(__file__).resolve().parent.parent
TMP = ROOT / "build_tmp"

BAGNET = "https://opendata.bagnet.ch/?r=/download&path="
REGIONS_URL = "https://www.priminfo.admin.ch/downloads/praemienregionen.xlsx"
INSURERS_URL = "https://www.bag.admin.ch/dam/de/sd-web/wKeV97535ICf/Zugelassene%20Krankenversicherer_1.1.2026.xlsx"


def bagnet_url(path: str) -> str:
    return BAGNET + urllib.parse.quote(base64.b64encode(path.encode()).decode())


def fetch(url: str, dest: Path) -> Path:
    if not dest.exists():
        print("download", url)
        with httpx.stream("GET", url, headers=UA, timeout=600, follow_redirects=True) as r:
            r.raise_for_status()
            with dest.open("wb") as f:
                for chunk in r.iter_bytes():
                    f.write(chunk)
    return dest


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", type=int, default=2026)
    args = ap.parse_args()
    TMP.mkdir(exist_ok=True)

    zpath = fetch(bagnet_url(f"/Praemien/Archiv_Praemien_{args.year}.zip"), TMP / f"praemien_{args.year}.zip")
    regions = fetch(REGIONS_URL, TMP / "praemienregionen.xlsx")
    insurers = fetch(INSURERS_URL, TMP / "insurers.xlsx")

    out = ROOT / "data" / f"premiums_{args.year}.sqlite"
    out.unlink(missing_ok=True)
    db = sqlite3.connect(out)
    db.executescript(
        """
        CREATE TABLE meta(key TEXT PRIMARY KEY, value TEXT);
        CREATE TABLE insurers(id INTEGER PRIMARY KEY, name TEXT);
        CREATE TABLE regions(bfs INTEGER, municipality TEXT, canton TEXT, region INTEGER, plz INTEGER, locality TEXT);
        CREATE TABLE premiums(
            insurer INTEGER, canton TEXT, region INTEGER, age_class TEXT, age_sub TEXT,
            accident INTEGER, tariff TEXT, tariff_type TEXT, tariff_name TEXT, franchise INTEGER, premium REAL);
        CREATE TABLE restricted(insurer INTEGER, canton TEXT, region INTEGER, tariff TEXT, bfs INTEGER);
        """
    )

    # Insurer names: first sheet is a compact list (number, name, locality).
    wb = openpyxl.load_workbook(insurers, read_only=True)
    names: dict[int, str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [c for c in row if c not in (None, "x")]
            if len(cells) >= 2 and isinstance(cells[0], int) and isinstance(cells[1], str):
                names.setdefault(cells[0], " ".join(cells[1].replace("-\n", "").split()))
    db.executemany("INSERT INTO insurers VALUES (?,?)", names.items())

    # Premium regions per municipality (BFS number is authoritative, PLZ only indicative).
    wb = openpyxl.load_workbook(regions, read_only=True)
    ws = wb["A_COM"]  # BFS-Nr, Kanton, Gemeinde, Region, Bezirk, PLZ, Ort
    rows = []
    for r in ws.iter_rows(values_only=True):
        bfs, canton, muni, region, _district, plz, loc = (r + (None,) * 7)[:7]
        if isinstance(bfs, int) and isinstance(region, int):
            rows.append((bfs, muni, canton, region, plz, loc))
    db.executemany("INSERT INTO regions VALUES (?,?,?,?,?,?)", rows)

    with zipfile.ZipFile(zpath) as z:
        with z.open("Prämien_CH.csv") as f:
            reader = csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig"))
            batch = []
            for r in reader:
                batch.append((
                    int(r["Versicherer"]), r["Kanton"], int(r["Region"].split("CH")[-1]),
                    r["Altersklasse"], r["Altersuntergruppe"] or None,
                    1 if r["Unfalleinschluss"] == "MIT-UNF" else 0,
                    r["Tarif"], r["Tariftyp"], r["Tarifbezeichnung"],
                    int(r["Franchise"].split("-")[1]), float(r["Prämie"]),
                ))
            db.executemany("INSERT INTO premiums VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)

        # Some alternative models (mostly HMO) are only sold in listed municipalities.
        with z.open("Einzugsgebiete.csv") as f:
            for r in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig"), delimiter=";"):
                if r["Eingeschränkt"] == "Y":
                    db.executemany("INSERT INTO restricted VALUES (?,?,?,?,?)", [
                        (int(r["Versicherer"]), r["Kanton"], int(r["Region"].split("CH")[-1]), r["Tarif"], int(b))
                        for b in r["Gemeinden-BFS"].split(",") if b.strip()
                    ])

    db.executescript(
        """
        CREATE INDEX ix_p ON premiums(canton, region, age_class, franchise, accident, premium);
        CREATE INDEX ix_r ON regions(bfs);
        """
    )
    db.executemany("INSERT INTO meta VALUES (?,?)", [
        ("year", str(args.year)),
        ("built", date.today().isoformat()),
        ("premium_source", bagnet_url(f"/Praemien/Archiv_Praemien_{args.year}.zip")),
        ("regions_source", REGIONS_URL),
        ("insurers_source", INSURERS_URL),
    ])
    db.commit()
    db.execute("VACUUM")
    print(out, out.stat().st_size // 1024, "KiB", len(names), "insurers", len(rows), "region rows")


if __name__ == "__main__":
    main()
